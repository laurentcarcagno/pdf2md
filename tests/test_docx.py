from io import BytesIO
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.shared import Inches

from pdf2md.docx_cli import main
from pdf2md.docx_convert import convert_docx, open_docx
from pdf2md.docx_numbering import format_number, parse_numbering
from pdf2md.errors import DocxOpenError
from pdf2md.models import Settings

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _png(color: tuple[int, int, int]) -> BytesIO:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 40), False)
    pix.set_rect(pix.irect, color)
    return BytesIO(pix.tobytes("png"))


@pytest.fixture()
def sample_docx(tmp_path: Path) -> Path:
    document = docx.Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    section = document.sections[0]
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "ACME Confidential"
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = "Page footer"

    document.add_paragraph("Table of Contents", style="TOC Heading")
    document.add_paragraph("1 Scope\t3", style="toc 1")
    document.add_paragraph("2 Terms\t4", style="toc 1")
    document.add_heading("Scope", level=1)
    document.add_paragraph("This document has a first paragraph.")
    document.add_paragraph("First step", style="List Number")
    document.add_paragraph("Second step", style="List Number")
    document.add_paragraph("Bullet item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    for index, text in enumerate(("Name", "Value", "Speed", "Fast")):
        table.cell(index // 2, index % 2).text = text
    document.add_picture(_png((200, 30, 30)), width=Inches(1))
    document.add_paragraph("Figure 1: A red box")
    document.add_page_break()
    document.add_heading("Terms", level=2)
    document.add_paragraph("Second page text.")
    path = tmp_path / "sample.docx"
    document.save(str(path))
    return path


def test_structure_is_converted(sample_docx: Path, tmp_path: Path) -> None:
    result = convert_docx(sample_docx, Settings(), images_dir=tmp_path / "assets")
    markdown = result.markdown
    assert "# Scope" in markdown and "## Terms" in markdown
    assert "This document has a first paragraph." in markdown
    assert "1. First step\n2. Second step" in markdown
    assert "- Bullet item" in markdown
    assert "| Name | Value |" in markdown and "| Speed | Fast |" in markdown
    assert "![Figure 1: A red box](assets/p0001_01.png)" in markdown


def test_headers_footers_and_toc_removed_by_default(sample_docx: Path) -> None:
    markdown = convert_docx(sample_docx, Settings()).markdown
    for unwanted in ("ACME Confidential", "Page footer", "Table of Contents", "1 Scope 3"):
        assert unwanted not in markdown


def test_keep_restores_elements_and_tags_them(sample_docx: Path) -> None:
    result = convert_docx(sample_docx, Settings(keep=frozenset({"toc", "header", "footer"})))
    assert "ACME Confidential" in result.markdown and "- 1 Scope 3" in result.markdown
    elements = {e.element for e in result.mapping if e.element}
    assert elements == {"toc", "header", "footer"}


def test_pages_are_estimated_from_page_breaks_and_selectable(sample_docx: Path) -> None:
    result = convert_docx(sample_docx, Settings())
    pages = {e.snippet: e.page_number for e in result.mapping}
    assert pages["Scope"] == 1 and pages["Terms"] == 2
    only_second = convert_docx(sample_docx, Settings(pages="2"))
    assert [e.snippet for e in only_second.mapping] == ["Terms", "Second page text."]
    with pytest.raises(ValueError):
        convert_docx(sample_docx, Settings(pages="9-"))


def test_image_files_are_written_only_with_images_dir(sample_docx: Path, tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assert not [b for b in convert_docx(sample_docx, Settings()).blocks if b.kind == "image"]
    convert_docx(sample_docx, Settings(), images_dir=assets)
    assert [p.name for p in assets.iterdir()] == ["p0001_01.png"]


def test_toc_field_and_list_of_figures_are_tagged(tmp_path: Path) -> None:
    document = docx.Document()
    body = document.element.body
    for instruction, text in (
        ('TOC \\o "1-3"', "Chapter 1 ..... 5"),
        ('TOC \\c "Figure"', "Figure 1 ..... 6"),
        ('TOC \\c "Table"', "Table 1 ..... 7"),
    ):
        body.insert(
            len(body) - 1,
            parse_xml(
                f"<w:p {W}><w:fldSimple w:instr='{instruction}'>"
                f"<w:r><w:t>{text}</w:t></w:r></w:fldSimple></w:p>"
            ),
        )
    document.add_paragraph("Real content.")
    path = tmp_path / "fields.docx"
    document.save(str(path))
    result = convert_docx(
        path,
        Settings(
            keep=frozenset({"toc", "lof", "lot"})
            and frozenset({"toc", "list_of_figures", "list_of_tables"})
        ),
    )
    assert [e.element for e in result.mapping] == [
        "toc",
        "list_of_figures",
        "list_of_tables",
        None,
    ]
    assert convert_docx(path, Settings()).markdown.strip() == "Real content."


def test_repeated_images_are_decoration(tmp_path: Path) -> None:
    document = docx.Document()
    logo = _png((0, 0, 200)).getvalue()
    for _ in range(3):
        document.add_picture(BytesIO(logo))
    document.add_picture(_png((200, 200, 0)))
    path = tmp_path / "logos.docx"
    document.save(str(path))
    result = convert_docx(path, Settings(), images_dir=tmp_path / "assets")
    assert sum(b.kind == "image" for b in result.blocks) == 1


def test_numbering_labels_and_nesting() -> None:
    xml = f"""<w:numbering {W}>
      <w:abstractNum w:abstractNumId="0">
        <w:lvl w:ilvl="0">
          <w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1."/>
        </w:lvl>
        <w:lvl w:ilvl="1">
          <w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/>
        </w:lvl>
        <w:lvl w:ilvl="2">
          <w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:lvlText w:val="(%3)"/>
        </w:lvl>
      </w:abstractNum>
      <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
    </w:numbering>""".encode()
    numbering = parse_numbering(xml)
    assert [numbering.label("1", 0), numbering.label("1", 1), numbering.label("1", 1)] == [
        "a.",
        "a.1",
        "a.2",
    ]
    assert numbering.label("1", 2) == "(i)"
    assert numbering.label("1", 0) == "b."
    assert numbering.label("1", 1) == "b.1"  # deeper counters restart
    assert numbering.label("0", 0) is None
    assert [format_number(n, "upperRoman") for n in (4, 9, 14)] == ["IV", "IX", "XIV"]


def test_lettered_items_are_separate_paragraphs_with_nested_list(tmp_path: Path) -> None:
    document = docx.Document()
    numbering_part = document.part.numbering_part.element
    numbering_part.append(
        parse_xml(
            f"""<w:abstractNum {W} w:abstractNumId="90">
              <w:lvl w:ilvl="0">
          <w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1."/>
        </w:lvl>
              <w:lvl w:ilvl="1">
          <w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%2."/>
        </w:lvl>
            </w:abstractNum>"""
        )
    )
    numbering_part.append(
        parse_xml(f'<w:num {W} w:numId="90"><w:abstractNumId w:val="90"/></w:num>')
    )

    def numbered(text: str, level: int) -> None:
        paragraph = document.add_paragraph(text)
        paragraph._p.get_or_add_pPr().append(
            parse_xml(f'<w:numPr {W}><w:ilvl w:val="{level}"/><w:numId w:val="90"/></w:numPr>')
        )

    numbered("First requirement", 0)
    document.add_paragraph("NOTE A note about the first requirement")
    numbered("Second requirement shall consist of:", 0)
    numbered("one", 1)
    numbered("two", 1)
    path = tmp_path / "req.docx"
    document.save(str(path))
    markdown = convert_docx(path, Settings()).markdown
    assert markdown == (
        "a. First requirement\n\nNOTE A note about the first requirement\n\n"
        "b. Second requirement shall consist of:\n\n1. one\n2. two\n"
    )


def test_open_errors(tmp_path: Path) -> None:
    with pytest.raises(DocxOpenError):
        open_docx(tmp_path / "missing.docx")
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(DocxOpenError):
        open_docx(bad)
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"x")
    with pytest.raises(DocxOpenError, match="legacy"):
        open_docx(legacy)


def test_cli_layout_and_exit_codes(sample_docx: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    assert main([str(sample_docx), "-o", str(out), "--no-progress"]) == 0
    folder = out / "sample"
    assert sorted(p.name for p in folder.iterdir()) == ["assets", "sample.json", "sample.md"]
    other = tmp_path / "other"
    assert (
        main([str(sample_docx), "-o", str(other), "--no-images", "--no-json", "--no-progress"]) == 0
    )
    assert [p.name for p in (other / "sample").iterdir()] == ["sample.md"]
    assert main([str(tmp_path / "nope.docx"), "-o", str(out), "--no-progress"]) == 1
